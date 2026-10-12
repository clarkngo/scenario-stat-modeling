"""Build a 16:9 PowerPoint deck from paper/content.json and results/results.json.

Usage:  python3 tools/build_pptx.py      ->  build/scenario-stat-modeling.pptx
"""
import json
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BG, BG_SOFT = RGBColor(0x0F, 0x11, 0x15), RGBColor(0x16, 0x19, 0x22)
FG, DIM, ACCENT, BORDER = (RGBColor(0xE8, 0xEA, 0xF0), RGBColor(0x9A, 0xA2, 0xB1),
                           RGBColor(0x5B, 0x8C, 0xFF), RGBColor(0x26, 0x2B, 0x38))
FONT = "Calibri"
W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.8)


def text(slide, s, left, top, width, height, size, color=FG, bold=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text, r.font.size, r.font.bold, r.font.color.rgb, r.font.name = s, Pt(size), bold, color, FONT
    return tf


def blank(prs, kicker=None, title=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    if kicker:
        text(slide, kicker.upper(), MARGIN, Inches(0.6), Inches(11), Inches(0.4), 13, ACCENT, True)
    if title:
        text(slide, title, MARGIN, Inches(1.0), Inches(11.7), Inches(0.9), 32, FG, True)
    return slide


def bullets(slide, items, top=Inches(2.2), width=Inches(11.5)):
    tf = slide.shapes.add_textbox(MARGIN, top, width, Inches(4.5)).text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        dot = p.add_run()
        dot.text, dot.font.color.rgb, dot.font.size = "●  ", ACCENT, Pt(14)
        r = p.add_run()
        r.text, r.font.size, r.font.color.rgb, r.font.name = item, Pt(22), FG, FONT


def table(slide, header, rows, top=Inches(2.2), width=Inches(11.7)):
    shape = slide.shapes.add_table(len(rows) + 1, len(header), MARGIN, top, width,
                                   Inches(0.45) * (len(rows) + 1))
    tbl = shape.table
    for i, row in enumerate([header] + rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BG_SOFT if i == 0 else BG
            cell.text = val
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size, run.font.name = Pt(14 if i == 0 else 16), FONT
            run.font.bold = i == 0
            run.font.color.rgb = ACCENT if i == 0 else FG


def build(out):
    with open(os.path.join(ROOT, "paper", "content.json")) as f:
        paper = json.load(f)
    with open(os.path.join(ROOT, "results", "results.json")) as f:
        results = json.load(f)
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    a = paper["author"]
    s = blank(prs, a["course"])
    text(s, paper["short_title"], MARGIN, Inches(2.2), Inches(11.7), Inches(1.2), 48, FG, True)
    text(s, "An AI-guided framework for deriving and revising algorithmic baselines from scratch",
         MARGIN, Inches(3.5), Inches(11), Inches(1), 22, DIM)
    text(s, f"{a['name']} — {a['program']}, {a['university']}", MARGIN, Inches(5.6),
         Inches(11), Inches(0.5), 16, DIM)

    s = blank(prs, "Framework", "Derive → Implement → Shock → Diagnose → Revise")
    bullets(s, ["Derive the estimator from a loss function, by hand",
                "Implement it in vectorized NumPy — no scikit-learn, no SciPy fitting",
                "Shock it with a realistic data failure; evaluate on clean data",
                "Diagnose which assumption broke; derive a targeted revision"])

    for sec in paper["sections"]:
        s = blank(prs, f"Section {sec['number']}", sec["title"])
        bullets(s, sec["slides"])
        if sec["id"] == "results":
            s = blank(prs, "Results", "Clean-test RMSE under the scenario shock ($1000s)")
            rows = [[f"{r['contamination']:.0%}", f"{r['ols_rmse']:.1f}", f"{r['hub_rmse']:.1f}",
                     f"{r['ols_berr']:.1f}", f"{r['hub_berr']:.1f}"] for r in results["rows"]]
            table(s, ["Shock", "OLS RMSE", "Huber RMSE", "OLS ‖β̂−β‖", "Huber ‖β̂−β‖"], rows)
            s = blank(prs, "Results", "OLS error grows with the shock; Huber stays near σ")
            s.shapes.add_picture(os.path.join(ROOT, "assets", "fig_sweep.png"),
                                 Inches(3.2), Inches(2.0), height=Inches(4.9))

    s = blank(prs, "Resources", "Paper, code, and live demo")
    bullets(s, ["Portal: https://clarkngo.github.io/scenario-stat-modeling/",
                "Slides: https://clarkngo.github.io/scenario-stat-modeling/presentations/ds623-scenario-stat-modeling.html",
                "Code: https://github.com/clarkngo/scenario-stat-modeling"])

    os.makedirs(os.path.dirname(out), exist_ok=True)
    prs.save(out)
    return out


if __name__ == "__main__":
    out = build(os.path.join(ROOT, "build", "scenario-stat-modeling.pptx"))
    print("wrote", os.path.relpath(out, ROOT))
