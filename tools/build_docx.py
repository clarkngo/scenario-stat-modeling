"""Build the paper as a Word document in the course-paper format (see README).

Layout: single-column title block and abstract, two-column body, numbered
ALL-CAPS section headings, APA-style "Table n"/"Figure n" labels, hanging-indent
references, and single-column appendices on a new page.

Usage:  python3 tools/build_docx.py      ->  build/scenario-stat-modeling.docx
"""
import json
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FONT, MONO, MATH = "Verdana", "Courier New", "Cambria Math"
COL_WIDTH = Inches(3.1)


def load():
    with open(os.path.join(ROOT, "paper", "content.json")) as f:
        paper = json.load(f)
    with open(os.path.join(ROOT, "results", "results.json")) as f:
        results = json.load(f)
    return paper, results


def results_table(results):
    header = ["Shock", "OLS RMSE", "Huber RMSE", "OLS ‖β̂−β‖", "Huber ‖β̂−β‖"]
    rows = [[f"{r['contamination']:.0%}", f"{r['ols_rmse']:.1f}", f"{r['hub_rmse']:.1f}",
             f"{r['ols_berr']:.1f}", f"{r['hub_berr']:.1f}"] for r in results["rows"]]
    return header, rows


def set_columns(section, num):
    cols = section._sectPr.xpath("./w:cols")
    col = cols[0] if cols else OxmlElement("w:cols")
    col.set(qn("w:num"), str(num))
    col.set(qn("w:space"), "432")  # 0.3 inch gutter
    if not cols:
        section._sectPr.append(col)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def para(doc, text="", size=10, bold=False, italic=False, align=None, before=0, after=8,
         font=FONT, keep_next=False):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.font.size, r.font.name, r.bold, r.italic = Pt(size), font, bold, italic
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.0
    pf.keep_with_next = keep_next
    if align is not None:
        p.alignment = align
    return p


class Builder:
    def __init__(self, paper, results):
        self.paper, self.results = paper, results
        self.doc = Document()
        self.n_table = self.n_fig = self.n_eq = 0
        st = self.doc.styles["Normal"]
        st.font.name, st.font.size = FONT, Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Inches(1))

    # --- front matter -------------------------------------------------------
    def title_block(self):
        p, a = self.paper, self.paper["author"]
        para(self.doc, p["title"], size=18, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
        for line in (a["name"], a["course"], a["program"], a["school"], a["university"], a["email"]):
            para(self.doc, line, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
        para(self.doc, "Abstract", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
             before=16, after=8)
        para(self.doc, p["abstract"], align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=10)
        kp = para(self.doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=10)
        r = kp.add_run("Keywords: ")
        r.bold = True
        kp.add_run(p["keywords"])

    # --- body blocks --------------------------------------------------------
    def heading(self, text):
        para(self.doc, text.upper(), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
             before=10, after=10, keep_next=True)

    def block(self, b):
        t = b["type"]
        if t == "p":
            para(self.doc, b["text"], align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif t == "h3":
            para(self.doc, b["text"], bold=True, before=4, after=8, keep_next=True)
        elif t == "list":
            for item in b["items"]:
                lp = self.doc.add_paragraph(item, style="List Bullet")
                lp.paragraph_format.space_after = Pt(4)
            para(self.doc, after=2)
        elif t == "eq":
            self.n_eq += 1
            # Segments separated by three spaces go on their own centered line;
            # the equation number sits flush right on the last line.
            parts = b["text"].split("   ")
            for k, part in enumerate(parts):
                last = k == len(parts) - 1
                ep = para(self.doc, after=8 if last else 2, keep_next=not last)
                ep.paragraph_format.tab_stops.add_tab_stop(Inches(1.55), WD_TAB_ALIGNMENT.CENTER)
                ep.paragraph_format.tab_stops.add_tab_stop(COL_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
                r = ep.add_run("\t" + part.strip())
                r.font.name, r.italic, r.font.size = MATH, True, Pt(9.5)
                if last:
                    ep.add_run(f"\t({self.n_eq})")
        elif t == "table":
            self.table(b)
        elif t == "figure":
            self.figure(b)

    def table(self, b):
        self.n_table += 1
        header, rows = (results_table(self.results) if b.get("source") == "results"
                        else (b["header"], b["rows"]))
        para(self.doc, f"Table {self.n_table}", bold=True, after=0, before=4, keep_next=True)
        para(self.doc, b["title"], italic=True, after=4, keep_next=True)
        tbl = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate([header] + rows):
            for j, val in enumerate(row):
                cell = tbl.cell(i, j)
                cell.text = ""
                r = cell.paragraphs[0].add_run(val)
                r.font.size, r.bold = Pt(7.5), i == 0
                if i == 0:
                    shade(cell, "F2F2F2")
        para(self.doc, b["caption"], italic=True, size=9, before=2, after=10)

    def figure(self, b):
        self.n_fig += 1
        para(self.doc, f"Figure {self.n_fig}", bold=True, after=0, before=4, keep_next=True)
        para(self.doc, b["title"], italic=True, after=4, keep_next=True)
        fp = para(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, keep_next=True)
        fp.add_run().add_picture(os.path.join(ROOT, b["src"]), width=COL_WIDTH)
        para(self.doc, b["caption"], italic=True, size=9, after=10)

    # --- back matter --------------------------------------------------------
    def references(self, number):
        self.heading(f"{number}. References")
        for ref in sorted(self.paper["references"]):
            rp = para(self.doc, ref, size=9, after=6)
            rp.paragraph_format.left_indent = Inches(0.3)
            rp.paragraph_format.first_line_indent = Inches(-0.3)

    def appendix(self):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        set_columns(sec, 1)
        for i, a in enumerate(self.paper["appendix"]):
            para(self.doc, a["title"].upper(), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                 before=12 if i else 0, after=10, keep_next=True)
            if "code" in a:
                with open(os.path.join(ROOT, a["code"])) as f:
                    for line in f.read().rstrip().split("\n"):
                        cp = para(self.doc, line or " ", size=7.5, font=MONO, after=0)
                        cp.runs[0].font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
            else:
                para(self.doc, a["text"], align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    def build(self, out):
        self.title_block()
        body = self.doc.add_section(WD_SECTION.CONTINUOUS)
        set_columns(body, 2)
        for s in self.paper["sections"]:
            self.heading(f"{s['number']}. {s['title']}")
            for b in s["blocks"]:
                self.block(b)
        self.references(len(self.paper["sections"]) + 1)
        self.appendix()
        os.makedirs(os.path.dirname(out), exist_ok=True)
        self.doc.save(out)
        return out


if __name__ == "__main__":
    paper, results = load()
    out = Builder(paper, results).build(os.path.join(ROOT, "build", "scenario-stat-modeling.docx"))
    print("wrote", os.path.relpath(out, ROOT))
